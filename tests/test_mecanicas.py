"""Pruebas del módulo de cargas mecánicas.

Los valores esperados son los que calcula la hoja 'SMC-A_15' del libro de
Excel de referencia, hasta el último decimal.
"""
from __future__ import annotations

import pytest

from app import mec_backup as B
from app import mecanicas as M
from app import parsing
from tests import fixtures_mec as fx

# Resultados de la hoja SMC-A_15 del Excel original.
ESPERADO = {
    "luz_viento": 55.0,
    "luz_peso": 102.0,
    "tension": {"1": 768.865864, "2": 263.086728, "3": 452.753904},
    "transversal": {"1": 95.5570825, "2": 72.3516713, "3": 52.1589060},
    "vertical": {"1": 62.96633495, "2": 33.82599679, "3": 56.43934729},
    "transversal_total": 411.1818249067565,
    "momento_total": 4765.11191652035,
    "momentos": {0.16: 3537.5231945708388, 1.74: 561.2298288888985, 3.29: 666.3588930606131},
}


@pytest.fixture(scope="module")
def dataset() -> M.MecDataset:
    ts = parsing.read_report(fx.tension_xlsx(), "tension.xlsx", "cable")
    ss = parsing.read_report(fx.span_xlsx(), "vanos.xlsx", "spans")
    return M.build_mec_dataset(ts, parsing.auto_map(ts, "cable"), ss, parsing.auto_map(ss, "spans"))


@pytest.fixture
def catalogo() -> dict[float, M.Conductor]:
    return {
        0.43: M.Conductor(0.43, "CTO1", 16.3068, 0.428477, 1),
        0.18: M.Conductor(0.18, "FO", 14.0, 0.181, 1),
        0.40: M.Conductor(0.40, "G", 9.144, 0.398413, 1),
    }


@pytest.fixture
def aisladores() -> dict[str, M.Aislador]:
    return {
        "Cadena 15 kV": M.Aislador("Cadena 15 kV", 115.0, 635.0, 3.4, 1, 15.0),
        "Solo ferretería": M.Aislador("Solo ferretería", 0.0, 0.0, 0.0, 0, 15.0),
    }


@pytest.fixture
def grupo() -> M.Grupo:
    return M.Grupo(
        nombre="SMC-A_15", estructuras=["6", "7"], casos=fx.CASOS_ELEGIDOS, n_postes=1,
        fs=1.1, nc=1, espesor_hielo_mm=0, n_cadenas=1, n_aisladores=1, alpha_deg=1,
        ht_m=15, t_servicio_kg=800, pv_kg_m2=80,
        lineas=[
            M.Linea("1", 0.43, fases=3, aislador="Cadena 15 kV", alturas_amarre=[0.16, 0.16, 0.16]),
            M.Linea("2", 0.18, fases=1, aislador="Solo ferretería", alturas_amarre=[3.29]),
            M.Linea("3", 0.40, fases=1, aislador="Solo ferretería", alturas_amarre=[1.74]),
        ],
    )


# --------------------------------------------------------------------------
# Lectura de los reportes
# --------------------------------------------------------------------------
def test_reports_are_recognised():
    ts = parsing.read_report(fx.tension_xlsx(), "t.xlsx", "cable")
    ss = parsing.read_report(fx.span_xlsx(), "v.xlsx", "spans")
    assert parsing.auto_map(ts, "cable").missing_required() == []
    assert parsing.auto_map(ss, "spans").missing_required() == []
    assert parsing.weight_span_columns(list(ss.columns)) == {
        1: "Weight Span LC# 1 (m)", 2: "Weight Span LC# 2 (m)", 3: "Weight Span LC# 3 (m)",
        4: "Weight Span LC# 4 (m)", 5: "Weight Span LC# 5 (m)",
    }


def test_weight_span_columns_are_matched_by_case_name_not_position(dataset):
    """El vano peso se toma del caso elegido, no de la primera columna."""
    assert dataset.weight_span["6"]["Viento Máximo"] == 102.0
    assert dataset.luz_peso(["6"], ["Viento Máximo"]) == 102.0
    assert dataset.luz_peso(["6"], ["Hielo"]) is None  # el hielo no tiene columna


def test_each_span_end_is_recorded_with_its_own_set(dataset):
    """El vano 7-8 sale del set 1 y llega al set 5."""
    en_8 = {(a.set_no, a.cable) for a in dataset.amarres if a.estructura == "8"}
    assert ("5", 0.43) in en_8
    assert ("1", 0.43) not in en_8
    en_7 = {(a.set_no, a.cable) for a in dataset.amarres if a.estructura == "7"}
    assert ("1", 0.43) in en_7


# --------------------------------------------------------------------------
# Opciones (set, cable)
# --------------------------------------------------------------------------
def test_options_pair_each_set_with_its_cable(dataset):
    opciones = dataset.opciones(["6", "7"], fx.CASOS_ELEGIDOS)
    assert [(o["set"], o["cable"]) for o in opciones] == [("1", 0.43), ("2", 0.18), ("3", 0.40)]
    assert opciones[0]["etiqueta"] == "Set 1 · 0.43 daN/m"


def test_one_set_carrying_two_cables_gives_two_options(dataset):
    """La estructura 8 tiene el set 5 con dos conductores distintos."""
    opciones = dataset.opciones(["8"], fx.CASOS_ELEGIDOS)
    assert [(o["set"], o["cable"]) for o in opciones] == [("5", 0.18), ("5", 0.43)]


def test_the_ice_case_is_excluded_when_not_selected(dataset):
    con_hielo = dataset.opciones(["6", "7"], fx.CASOS_ELEGIDOS + ["Hielo"])
    cables = {o["cable"] for o in con_hielo}
    assert 1.17 in cables  # con hielo aparece la variante
    assert 1.17 not in {o["cable"] for o in dataset.opciones(["6", "7"], fx.CASOS_ELEGIDOS)}


# --------------------------------------------------------------------------
# Cálculo: debe reproducir el Excel
# --------------------------------------------------------------------------
def test_spans_match_the_reference_sheet(dataset):
    assert dataset.luz_viento(["6", "7"]) == ESPERADO["luz_viento"]
    assert dataset.luz_peso(["6", "7"], fx.CASOS_ELEGIDOS) == ESPERADO["luz_peso"]


def test_line_results_match_the_reference_sheet(dataset, grupo, catalogo, aisladores):
    resultado = M.evaluar(dataset, grupo, catalogo, aisladores)
    assert [l.avisos for l in resultado.lineas] == [[], [], []]
    for linea in resultado.lineas:
        assert linea.tension_kg == pytest.approx(ESPERADO["tension"][linea.set_no], abs=1e-6)
        assert linea.carga_transversal == pytest.approx(ESPERADO["transversal"][linea.set_no], abs=1e-6)
        assert linea.carga_vertical == pytest.approx(ESPERADO["vertical"][linea.set_no], abs=1e-6)


def test_totals_and_moments_match_the_reference_sheet(dataset, grupo, catalogo, aisladores):
    r = M.evaluar(dataset, grupo, catalogo, aisladores)
    assert r.transversal_calculada == pytest.approx(ESPERADO["transversal_total"], abs=1e-9)
    assert r.transversal_admisible == 800
    assert r.momento_calculado == pytest.approx(ESPERADO["momento_total"], abs=1e-9)
    assert r.momento_admisible == 10000
    assert r.altura_efectiva == 12.5
    assert {m.altura_amarre: m.momento for m in r.momentos} == pytest.approx(ESPERADO["momentos"], abs=1e-9)
    assert r.cumple


def test_more_poles_raise_the_allowables(dataset, grupo, catalogo, aisladores):
    grupo.n_postes = 3
    r = M.evaluar(dataset, grupo, catalogo, aisladores)
    assert r.transversal_admisible == 2400
    assert r.momento_admisible == 30000


def test_a_line_that_does_not_exist_is_reported(dataset, grupo, catalogo, aisladores):
    grupo.lineas.append(M.Linea("9", 0.43, fases=1))
    r = M.evaluar(dataset, grupo, catalogo, aisladores)
    ultima = r.lineas[-1]
    assert ultima.tension_kg is None and ultima.carga_transversal is None
    assert "no aparece en esas estructuras" in ultima.avisos[0]


def test_a_missing_diameter_is_reported(dataset, grupo, catalogo, aisladores):
    catalogo[0.43] = M.Conductor(0.43, "CTO1", 0.0, 0.428477, 1)
    r = M.evaluar(dataset, grupo, catalogo, aisladores)
    assert any("diámetro" in a for a in r.lineas[0].avisos)


def test_the_tension_condition_can_be_changed():
    ts = parsing.read_report(fx.tension_xlsx(), "t.xlsx", "cable")
    ss = parsing.read_report(fx.span_xlsx(), "v.xlsx", "spans")
    inicial = M.build_mec_dataset(ts, parsing.auto_map(ts, "cable"), ss,
                                  parsing.auto_map(ss, "spans"), condicion="inicial")
    # La condición inicial del fixture es un 5% mayor que la de fluencia.
    esperado = ESPERADO["tension"]["1"] * 1.05
    assert inicial.tension_max(["6", "7"], "1", 0.43, fx.CASOS_ELEGIDOS) == pytest.approx(esperado, rel=1e-6)


# --------------------------------------------------------------------------
# Respaldo
# --------------------------------------------------------------------------
@pytest.fixture
def proyecto(grupo, catalogo, aisladores) -> B.Proyecto:
    otro = M.Grupo(
        nombre="Set repetido", estructuras=["8"], casos=fx.CASOS_ELEGIDOS, n_postes=2,
        lineas=[M.Linea("5", 0.43, fases=1, aislador="Cadena 15 kV", alturas_amarre=[0.5]),
                M.Linea("5", 0.18, fases=1, aislador="Solo ferretería", alturas_amarre=[1.5])],
    )
    return B.Proyecto(nombre="Proyecto de prueba", condicion="creep",
                      conductores=catalogo, aisladores=aisladores, grupos=[grupo, otro])


def test_backup_round_trip_reproduces_every_result(dataset, proyecto):
    antes = {g.nombre: M.evaluar(dataset, g, proyecto.conductores, proyecto.aisladores)
             for g in proyecto.grupos}
    blob = B.escribir_respaldo(proyecto, antes)

    recuperado = B.leer_respaldo(blob)
    assert recuperado.nombre == "Proyecto de prueba"
    assert recuperado.condicion == "creep"
    assert [g.nombre for g in recuperado.grupos] == ["SMC-A_15", "Set repetido"]

    despues = {g.nombre: M.evaluar(dataset, g, recuperado.conductores, recuperado.aisladores)
               for g in recuperado.grupos}
    for nombre, a in antes.items():
        d = despues[nombre]
        assert d.transversal_calculada == pytest.approx(a.transversal_calculada, abs=1e-9)
        assert d.momento_calculado == pytest.approx(a.momento_calculado, abs=1e-9)
        assert d.vertical_total == pytest.approx(a.vertical_total, abs=1e-9)
        assert [(l.set_no, l.cable, l.tension_kg) for l in d.lineas] == \
               [(l.set_no, l.cable, l.tension_kg) for l in a.lineas]


def test_backup_keeps_two_lines_with_the_same_set(dataset, proyecto):
    blob = B.escribir_respaldo(proyecto, {})
    grupo = B.leer_respaldo(blob).grupos[1]
    assert [(l.set_no, l.cable) for l in grupo.lineas] == [("5", 0.43), ("5", 0.18)]


def test_backup_keeps_the_conductor_catalogue(dataset, proyecto):
    recuperado = B.leer_respaldo(B.escribir_respaldo(proyecto, {}))
    cto = recuperado.conductores[0.43]
    assert (cto.nombre, cto.diametro_mm, cto.peso_dan_m) == ("CTO1", 16.3068, 0.428477)
    assert cto.peso_kg_m == pytest.approx(0.436924852, abs=1e-9)


def test_a_backup_is_told_apart_from_a_report(proyecto):
    assert B.es_respaldo(B.escribir_respaldo(proyecto, {}))
    assert not B.es_respaldo(fx.tension_xlsx())
    assert not B.es_respaldo(b"no soy un xlsx")


def test_reading_something_that_is_not_a_backup_explains_why():
    with pytest.raises(ValueError, match="no es un respaldo válido"):
        B.leer_respaldo(fx.tension_xlsx())


def test_backup_carries_readable_result_sheets(dataset, proyecto):
    import io

    import openpyxl

    resultados = {g.nombre: M.evaluar(dataset, g, proyecto.conductores, proyecto.aisladores)
                  for g in proyecto.grupos}
    libro = openpyxl.load_workbook(io.BytesIO(B.escribir_respaldo(proyecto, resultados)))
    assert libro.sheetnames[0] == "Resumen"
    assert "SMC-A_15" in libro.sheetnames
    resumen = [c.value for c in libro["Resumen"][2]]
    assert resumen[0] == "SMC-A_15"
    assert resumen[1] == pytest.approx(411.18, abs=0.01)
    assert resumen[-1] == "CUMPLE"


# --------------------------------------------------------------------------
# Datos del poste que trae el reporte
# --------------------------------------------------------------------------
def test_pole_data_is_read_from_the_lookup_block(dataset):
    """El reporte repite 'Structure File Name' con altura y cargas de ensayo."""
    poste = dataset.postes["6"]
    assert poste.archivo == fx.POSTE_15
    assert (poste.altura_m, poste.transversal_kg, poste.longitudinal_kg) == (15.0, 1600.0, 480.0)
    assert dataset.postes["8"].altura_m == 16.5


def test_pole_data_for_a_uniform_group(dataset):
    datos = dataset.datos_poste(["6", "7"])
    assert datos["altura_m"] == 15.0
    assert datos["transversal_kg"] == 1600.0
    assert datos["archivos"] == [fx.POSTE_15]
    assert datos["avisos"] == []


def test_a_mixed_group_warns_and_takes_the_worst_case(dataset):
    datos = dataset.datos_poste(["7", "8"])
    assert datos["transversal_kg"] == 1200.0  # la menor de las dos
    assert any("alturas distintas" in a for a in datos["avisos"])
    assert any("cargas transversales" in a for a in datos["avisos"])


def test_structures_without_pole_data_are_simply_absent(dataset):
    assert dataset.datos_poste(["999"]) == {
        "archivos": [], "altura_m": None, "transversal_kg": None,
        "longitudinal_kg": None, "avisos": [],
    }


def test_backup_keeps_the_insulator_catalogue(dataset, proyecto):
    recuperado = B.leer_respaldo(B.escribir_respaldo(proyecto, {}))
    cadena = recuperado.aisladores["Cadena 15 kV"]
    assert (cadena.diametro_mm, cadena.longitud_mm, cadena.peso_kg) == (115.0, 635.0, 3.4)
    assert (cadena.n_aisladores, cadena.peso_ferreteria_kg) == (1, 15.0)
    assert recuperado.grupos[0].lineas[0].aislador == "Cadena 15 kV"


def test_a_line_pointing_at_a_missing_insulator_is_reported(dataset, grupo, catalogo):
    r = M.evaluar(dataset, grupo, catalogo, {})   # catálogo de aisladores vacío
    assert any("no está en el catálogo" in a for a in r.lineas[0].avisos)


def test_an_old_backup_moves_the_hardware_into_the_catalogue(dataset, grupo, catalogo):
    """Un respaldo v1 llevaba el herraje en cada línea; se convierte en catálogo."""
    import io

    import openpyxl

    proyecto = B.Proyecto(nombre="Antiguo", conductores=catalogo, grupos=[grupo])
    libro = openpyxl.load_workbook(io.BytesIO(B.escribir_respaldo(proyecto, {})))
    libro[B.HOJA_PROYECTO]["B2"] = 1                      # versión 1
    del libro[B.HOJA_AISLADORES]
    hoja = libro[B.HOJA_LINEAS]
    libro.remove(hoja)
    hoja = libro.create_sheet(B.HOJA_LINEAS)
    hoja.append(B.CAB_LINEAS_V1)
    hoja.append(["SMC-A_15", "1", 0.43, 3, 1, 115, 635, 3.4, 1, 15, "0.16, 0.16, 0.16"])
    hoja.append(["SMC-A_15", "2", 0.18, 1, 1, 0, 0, 0, 0, 15, "3.29"])
    hoja.append(["SMC-A_15", "3", 0.40, 1, 1, 0, 0, 0, 0, 15, "1.74"])
    buffer = io.BytesIO()
    libro.save(buffer)

    recuperado = B.leer_respaldo(buffer.getvalue())
    # Dos cadenas distintas: la del CTO y la que solo lleva ferretería.
    assert len(recuperado.aisladores) == 2
    lineas = recuperado.grupos[0].lineas
    assert lineas[1].aislador == lineas[2].aislador != lineas[0].aislador

    r = M.evaluar(dataset, recuperado.grupos[0], recuperado.conductores, recuperado.aisladores)
    assert r.transversal_calculada == pytest.approx(ESPERADO["transversal_total"], abs=1e-9)
    assert r.momento_calculado == pytest.approx(ESPERADO["momento_total"], abs=1e-9)
