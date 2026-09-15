"""Pruebas HTTP del módulo de cargas mecánicas."""
from __future__ import annotations

import importlib
import io
import os

import openpyxl
import pytest

from tests import fixtures_mec as fx

PASSWORD = "clave-de-prueba"


@pytest.fixture(scope="module")
def client():
    os.environ["APP_PASSWORD"] = PASSWORD
    os.environ["SECRET_KEY"] = "secreto-para-tests"
    os.environ["COOKIE_SECURE"] = "0"
    from fastapi.testclient import TestClient

    from app import auth, config, main, mec_routes, store

    importlib.reload(config)
    importlib.reload(auth)
    importlib.reload(store)
    importlib.reload(mec_routes)
    importlib.reload(main)
    cliente = TestClient(main.app)
    cliente.post("/login", data={"password": PASSWORD}, follow_redirects=False)
    return cliente


def _archivos(incluir_respaldo: bytes | None = None):
    archivos = [
        ("archivos", ("tension.xlsx", fx.tension_xlsx(), "application/vnd.ms-excel")),
        ("archivos", ("vanos.xlsx", fx.span_xlsx(), "application/vnd.ms-excel")),
    ]
    if incluir_respaldo:
        archivos.append(("archivos", ("respaldo.xlsx", incluir_respaldo, "application/vnd.ms-excel")))
    return archivos


@pytest.fixture(scope="module")
def job(client):
    r = client.post("/api/mec/upload", files=_archivos())
    assert r.status_code == 200, r.text
    return r.json()


# --------------------------------------------------------------------------
# Páginas
# --------------------------------------------------------------------------
def test_home_offers_both_tools(client):
    cuerpo = client.get("/").text
    assert "Tablas de tensado" in cuerpo and "Cargas mecánicas" in cuerpo
    assert 'href="/tensado"' in cuerpo and 'href="/cargas"' in cuerpo


def test_each_tool_has_its_own_page(client):
    assert "Cargar reportes" in client.get("/tensado").text
    assert "Cargas mecánicas" in client.get("/cargas").text


def test_the_tool_pages_need_a_session(client):
    anonimo = client.__class__(client.app)
    for ruta in ("/", "/tensado", "/cargas"):
        assert "Clave de acceso" in anonimo.get(ruta).text


def test_the_api_needs_a_session(client):
    anonimo = client.__class__(client.app)
    assert anonimo.post("/api/mec/evaluar", json={"job_id": "x", "proyecto": {}}).status_code == 401


# --------------------------------------------------------------------------
# Carga de archivos
# --------------------------------------------------------------------------
def test_upload_identifies_both_reports(job):
    assert set(job["archivos"]) == {"tension", "spans"}
    assert job["sin_reconocer"] == []
    assert [e["key"] for e in job["estructuras"]] == ["5", "6", "7", "8", "9"]
    assert set(fx.CASOS_ELEGIDOS) <= set(job["casos"])
    assert job["proyecto"] is None


def test_upload_without_the_span_report_says_what_is_missing(client):
    r = client.post("/api/mec/upload", files=[
        ("archivos", ("tension.xlsx", fx.tension_xlsx(), "application/vnd.ms-excel")),
    ])
    assert r.status_code == 400
    assert "Reporte de vanos" in r.json()["detail"]


# --------------------------------------------------------------------------
# Opciones (set, cable)
# --------------------------------------------------------------------------
def test_options_are_pairs_of_set_and_cable(client, job):
    r = client.post("/api/mec/opciones", json={
        "job_id": job["job_id"], "estructuras": ["6", "7"], "casos": fx.CASOS_ELEGIDOS,
    })
    opciones = r.json()["opciones"]
    assert [(o["set"], o["cable"]) for o in opciones] == [("1", 0.43), ("2", 0.18), ("3", 0.40)]
    assert opciones[0]["etiqueta"] == "Set 1 · 0.43 daN/m"


def test_options_show_the_same_set_twice_when_it_carries_two_cables(client, job):
    r = client.post("/api/mec/opciones", json={
        "job_id": job["job_id"], "estructuras": ["8"], "casos": fx.CASOS_ELEGIDOS,
    })
    assert [(o["set"], o["cable"]) for o in r.json()["opciones"]] == [("5", 0.18), ("5", 0.43)]


# --------------------------------------------------------------------------
# Cálculo
# --------------------------------------------------------------------------
def _proyecto():
    return {
        "nombre": "Proyecto web", "condicion": "creep",
        "grupos": [{
            "nombre": "SMC-A_15", "estructuras": ["6", "7"], "casos": fx.CASOS_ELEGIDOS,
            "n_postes": 1, "fs": 1.1, "nc": 1, "espesor_hielo_mm": 0, "n_cadenas": 1,
            "n_aisladores": 1, "alpha_deg": 1, "ht_m": 15, "t_servicio_kg": 800, "pv_kg_m2": 80,
            "lineas": [
                {"set_no": "1", "cable": 0.43, "nombre": "CTO1", "fases": 3,
                 "diametro_conductor_mm": 16.3068, "diam_aislador_mm": 115, "long_aislador_mm": 635,
                 "n_conductores": 1, "peso_conductor_kg_m": 1.019716 * 0.428477,
                 "peso_aislador_kg": 3.4, "n_aisladores": 1, "peso_ferreteria_kg": 15,
                 "alturas_amarre": [0.16, 0.16, 0.16]},
                {"set_no": "2", "cable": 0.18, "nombre": "FO", "fases": 1,
                 "diametro_conductor_mm": 14.0, "n_conductores": 1,
                 "peso_conductor_kg_m": 1.019716 * 0.181, "peso_ferreteria_kg": 15,
                 "alturas_amarre": [3.29]},
                {"set_no": "3", "cable": 0.40, "nombre": "G", "fases": 1,
                 "diametro_conductor_mm": 9.144, "n_conductores": 1,
                 "peso_conductor_kg_m": 1.019716 * 0.398413, "peso_ferreteria_kg": 15,
                 "alturas_amarre": [1.74]},
            ],
        }],
    }


def test_evaluate_reproduces_the_reference_sheet(client, job):
    r = client.post("/api/mec/evaluar", json={"job_id": job["job_id"], "proyecto": _proyecto()})
    resultado = r.json()["resultados"][0]
    assert resultado["transversal_calculada"] == pytest.approx(411.1818249067565, abs=1e-9)
    assert resultado["momento_calculado"] == pytest.approx(4765.11191652035, abs=1e-9)
    assert resultado["transversal_admisible"] == 800
    assert resultado["cumple"] is True
    assert [round(l["tension_kg"], 6) for l in resultado["lineas"]] == [768.865864, 263.086728, 452.753904]


def test_phase_heights_are_kept_per_phase(client, job):
    """Las fases de un mismo set pueden ir a alturas distintas."""
    proyecto = _proyecto()
    proyecto["grupos"][0]["lineas"][0]["alturas_amarre"] = [0.6, 1.15, 2.15]
    r = client.post("/api/mec/evaluar", json={"job_id": job["job_id"], "proyecto": proyecto})
    momentos = r.json()["resultados"][0]["momentos"]
    assert [m["altura_amarre"] for m in momentos] == [0.6, 1.15, 1.74, 2.15, 3.29]


def test_several_sheets_are_evaluated_at_once(client, job):
    proyecto = _proyecto()
    proyecto["grupos"].append({
        "nombre": "Set repetido", "estructuras": ["8"], "casos": fx.CASOS_ELEGIDOS, "n_postes": 2,
        "lineas": [{"set_no": "5", "cable": 0.43, "fases": 1, "alturas_amarre": [0.5]},
                   {"set_no": "5", "cable": 0.18, "fases": 1, "alturas_amarre": [1.5]}],
    })
    resultados = client.post("/api/mec/evaluar",
                             json={"job_id": job["job_id"], "proyecto": proyecto}).json()["resultados"]
    assert [r["nombre"] for r in resultados] == ["SMC-A_15", "Set repetido"]
    assert [(l["set_no"], l["cable"]) for l in resultados[1]["lineas"]] == [("5", 0.43), ("5", 0.18)]
    assert resultados[1]["transversal_admisible"] == 1600


def test_an_expired_job_gives_a_clear_message(client):
    r = client.post("/api/mec/evaluar", json={"job_id": "inexistente", "proyecto": _proyecto()})
    assert r.status_code == 404 and "venció" in r.json()["detail"]


# --------------------------------------------------------------------------
# Respaldo
# --------------------------------------------------------------------------
def test_backup_downloads_as_a_spreadsheet(client, job):
    r = client.post("/api/mec/respaldo", json={"job_id": job["job_id"], "proyecto": _proyecto()})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment; filename=" in r.headers["content-disposition"]
    libro = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Resumen" in libro.sheetnames and "Respaldo_grupos" in libro.sheetnames


def test_uploading_the_backup_with_the_reports_restores_the_project(client, job):
    blob = client.post("/api/mec/respaldo",
                       json={"job_id": job["job_id"], "proyecto": _proyecto()}).content

    r = client.post("/api/mec/upload", files=_archivos(incluir_respaldo=blob))
    assert r.status_code == 200, r.text
    datos = r.json()
    assert "respaldo" in datos["archivos"]
    proyecto = datos["proyecto"]
    assert proyecto is not None
    assert proyecto["nombre"] == "Proyecto web"
    assert [g["nombre"] for g in proyecto["grupos"]] == ["SMC-A_15"]

    # Y los resultados vuelven a salir iguales.
    resultado = client.post("/api/mec/evaluar",
                            json={"job_id": datos["job_id"], "proyecto": proyecto}).json()["resultados"][0]
    assert resultado["transversal_calculada"] == pytest.approx(411.1818249067565, abs=1e-9)
    assert resultado["momento_calculado"] == pytest.approx(4765.11191652035, abs=1e-9)


def test_the_restored_project_keeps_the_phase_heights(client, job):
    proyecto = _proyecto()
    proyecto["grupos"][0]["lineas"][0]["alturas_amarre"] = [0.6, 1.15, 2.15]
    blob = client.post("/api/mec/respaldo", json={"job_id": job["job_id"], "proyecto": proyecto}).content
    recuperado = client.post("/api/mec/upload", files=_archivos(incluir_respaldo=blob)).json()["proyecto"]
    assert recuperado["grupos"][0]["lineas"][0]["alturas_amarre"] == [0.6, 1.15, 2.15]


def test_options_also_return_the_pole_data(client, job):
    r = client.post("/api/mec/opciones", json={
        "job_id": job["job_id"], "estructuras": ["6", "7"], "casos": fx.CASOS_ELEGIDOS,
    })
    poste = r.json()["poste"]
    assert poste["altura_m"] == 15.0
    assert poste["transversal_kg"] == 1600.0
    assert poste["archivos"] == [fx.POSTE_15]


def test_upload_lists_the_pole_height_per_structure(job):
    por_clave = {e["key"]: e for e in job["estructuras"]}
    assert por_clave["6"]["altura_m"] == 15.0
    assert por_clave["8"]["transversal_kg"] == 1200.0


def test_the_restored_project_keeps_every_line_field(client, job):
    blob = client.post("/api/mec/respaldo",
                       json={"job_id": job["job_id"], "proyecto": _proyecto()}).content
    proyecto = client.post("/api/mec/upload", files=_archivos(incluir_respaldo=blob)).json()["proyecto"]
    linea = proyecto["grupos"][0]["lineas"][0]
    assert linea["nombre"] == "CTO1"
    assert linea["diametro_conductor_mm"] == 16.3068
    assert (linea["diam_aislador_mm"], linea["long_aislador_mm"]) == (115.0, 635.0)
    assert linea["peso_conductor_kg_m"] == pytest.approx(0.436924852, abs=1e-9)
    assert linea["alturas_amarre"] == [0.16, 0.16, 0.16]
