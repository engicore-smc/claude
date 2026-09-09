# -*- coding: utf-8 -*-
"""Convierte el .xls homologado a .xlsx (openpyxl no lee BIFF/.xls)."""
import sys
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

HOR = {0: 'general', 1: 'left', 2: 'center', 3: 'right', 4: 'fill',
       5: 'justify', 6: 'centerContinuous', 7: 'distributed'}
VER = {0: 'top', 1: 'center', 2: 'bottom', 3: 'justify', 4: 'distributed'}


def convert(inp, outp):
    rb = xlrd.open_workbook(inp, formatting_info=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name in rb.sheet_names():
        sh = rb.sheet_by_name(name)
        ws = wb.create_sheet(name)
        for c, ci in sh.colinfo_map.items():
            ws.column_dimensions[get_column_letter(c + 1)].width = ci.width / 256.0
        for r in range(sh.nrows):
            if r in sh.rowinfo_map and sh.rowinfo_map[r].height_mismatch:
                ws.row_dimensions[r + 1].height = sh.rowinfo_map[r].height / 20.0
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v in ('', None):
                    continue
                cell = ws.cell(row=r + 1, column=c + 1, value=v)
                xf = rb.xf_list[sh.cell_xf_index(r, c)]
                f = rb.font_list[xf.font_index]
                cell.font = Font(name=f.name, size=f.height / 20.0,
                                 bold=f.weight >= 700, italic=bool(f.italic),
                                 underline='single' if f.underline_type else None)
                cell.alignment = Alignment(
                    horizontal=HOR.get(xf.alignment.hor_align),
                    vertical=VER.get(xf.alignment.vert_align),
                    wrap_text=bool(xf.alignment.text_wrapped))
                fmt = rb.format_map[xf.format_key].format_str
                if fmt and fmt != 'General':
                    cell.number_format = fmt
        for r1, r2, c1, c2 in sh.merged_cells:
            r2 = min(r2, sh.nrows)
            if r2 - 1 > r1 or c2 - 1 > c1:
                ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1,
                               end_row=r2, end_column=c2)
    wb.save(outp)
    print('OK ->', outp)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
